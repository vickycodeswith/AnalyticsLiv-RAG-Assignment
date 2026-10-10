from pathlib import Path
import sqlite3, random, textwrap
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from docx import Document
from docx.shared import Inches

ROOT = Path(__file__).parent
DATA = ROOT / 'data'
DOCS = DATA / 'documents'
DOCS.mkdir(parents=True, exist_ok=True)
random.seed(42)

# ---------- Structured sample data ----------
db_path = DATA / 'novatech.db'
if db_path.exists(): db_path.unlink()
conn = sqlite3.connect(db_path)
conn.execute('PRAGMA foreign_keys = ON')
conn.executescript('''
CREATE TABLE projects (
    project_id INTEGER PRIMARY KEY,
    project_name TEXT NOT NULL,
    business_unit TEXT NOT NULL,
    budget_usd INTEGER NOT NULL,
    status TEXT NOT NULL,
    start_date TEXT NOT NULL,
    target_date TEXT NOT NULL
);
CREATE TABLE employees (
    employee_id INTEGER PRIMARY KEY,
    full_name TEXT NOT NULL,
    email TEXT UNIQUE NOT NULL,
    role TEXT NOT NULL,
    location TEXT NOT NULL,
    project_id INTEGER NOT NULL,
    hire_date TEXT NOT NULL,
    FOREIGN KEY (project_id) REFERENCES projects(project_id)
);
CREATE TABLE support_tickets (
    ticket_id INTEGER PRIMARY KEY,
    employee_id INTEGER NOT NULL,
    category TEXT NOT NULL,
    priority TEXT NOT NULL,
    status TEXT NOT NULL,
    created_date TEXT NOT NULL,
    resolution_hours REAL,
    summary TEXT NOT NULL,
    FOREIGN KEY (employee_id) REFERENCES employees(employee_id)
);
''')

units = ['Customer Analytics','Platform Engineering','Data Operations','People Systems','Retail Insights','Cloud Services','Security Operations','Product Research','Finance Automation','Customer Experience']
project_adjs = ['Atlas','Beacon','Cedar','Delta','Evergreen','Frontier','Gateway','Harbor','Indigo','Juniper','Keystone','Lighthouse','Meridian','Northstar','Orchid','Pioneer','Quartz','Riverstone','Summit','Trailblazer']
project_nouns = ['Insights','Migration','Portal','Forecast','Assist','Workflow','Quality','Signals','Modernization','Workspace']
statuses = ['Active','Planning','On Hold','Completed']
projects=[]
for i in range(1,61):
    name=f'{project_adjs[(i-1)%len(project_adjs)]} {project_nouns[(i*3)%len(project_nouns)]} {i:02d}'
    unit=units[(i-1)%len(units)]
    budget=random.randrange(25000,260000,5000)
    status=statuses[(i-1)%len(statuses)]
    month=((i-1)%12)+1
    projects.append((i,name,unit,budget,status,f'2025-{month:02d}-01',f'2026-{((month+2-1)%12)+1:02d}-28'))
conn.executemany('INSERT INTO projects VALUES (?,?,?,?,?,?,?)',projects)
first_names=['Aarav','Aisha','Arjun','Diya','Ishaan','Kavya','Mira','Neel','Priya','Rohan','Sara','Vihaan','Anaya','Dev','Ira','Kabir','Meera','Reyansh','Tara','Zoya']
last_names=['Shah','Patel','Mehta','Rao','Desai','Kapoor','Joshi','Nair','Iyer','Verma','Singh','Malhotra','Trivedi','Khan','Kulkarni']
roles=['Data Analyst','Software Engineer','Product Manager','QA Engineer','Business Analyst','ML Engineer','Support Specialist','DevOps Engineer','UX Researcher','Data Engineer']
locations=['Ahmedabad','Gandhinagar','Pune','Bengaluru','Hyderabad','Remote']
employees=[]
for i in range(1,101):
    name=f'{first_names[(i*7)%len(first_names)]} {last_names[(i*11)%len(last_names)]}'
    employees.append((i,name,f'employee{i:03d}@novatech.example',roles[(i-1)%len(roles)],locations[(i-1)%len(locations)],((i-1)%60)+1,f'202{(i%5)+1}-{((i-1)%12)+1:02d}-{((i*3)%27)+1:02d}'))
conn.executemany('INSERT INTO employees VALUES (?,?,?,?,?,?,?)',employees)
categories=['Access Request','Laptop Support','Data Quality','VPN','Software Installation','Account Issue','Reporting','Email','Cloud Access','Other']
priorities=['Low','Medium','High','Critical']
ticket_statuses=['Open','In Progress','Resolved','Closed']
summaries={
'Access Request':'Requested access to a shared analytics workspace',
'Laptop Support':'Laptop performance issue during a scheduled update',
'Data Quality':'Reported a mismatch between dashboard and source totals',
'VPN':'Could not connect to the company VPN from a remote network',
'Software Installation':'Requested installation of an approved development tool',
'Account Issue':'Multi-factor authentication reset request',
'Reporting':'Asked for help refreshing a weekly report',
'Email':'Mail client stopped syncing a shared inbox',
'Cloud Access':'Requested read-only access to a cloud project',
'Other':'General support question about internal tools'}
tickets=[]
for i in range(1,151):
    emp=((i*13-1)%100)+1
    cat=categories[(i*3)%len(categories)]
    priority=priorities[(i*7)%len(priorities)]
    status=ticket_statuses[(i*5)%len(ticket_statuses)]
    month=((i-1)%12)+1
    day=((i*5)%27)+1
    hours=None if status in ('Open','In Progress') else round(random.uniform(0.5,72),1)
    tickets.append((i,emp,cat,priority,status,f'2026-{month:02d}-{day:02d}',hours,summaries[cat]))
conn.executemany('INSERT INTO support_tickets VALUES (?,?,?,?,?,?,?,?)',tickets)
conn.commit()

# Export reproducible schema and data dump
schema = '''-- NovaTech Solutions sample database. Fictional data created for the AnalyticsLiv assignment.
PRAGMA foreign_keys = ON;
CREATE TABLE projects (project_id INTEGER PRIMARY KEY, project_name TEXT NOT NULL, business_unit TEXT NOT NULL, budget_usd INTEGER NOT NULL, status TEXT NOT NULL, start_date TEXT NOT NULL, target_date TEXT NOT NULL);
CREATE TABLE employees (employee_id INTEGER PRIMARY KEY, full_name TEXT NOT NULL, email TEXT UNIQUE NOT NULL, role TEXT NOT NULL, location TEXT NOT NULL, project_id INTEGER NOT NULL, hire_date TEXT NOT NULL, FOREIGN KEY (project_id) REFERENCES projects(project_id));
CREATE TABLE support_tickets (ticket_id INTEGER PRIMARY KEY, employee_id INTEGER NOT NULL, category TEXT NOT NULL, priority TEXT NOT NULL, status TEXT NOT NULL, created_date TEXT NOT NULL, resolution_hours REAL, summary TEXT NOT NULL, FOREIGN KEY (employee_id) REFERENCES employees(employee_id));
'''
(DATA/'schema.sql').write_text(schema)
with (DATA/'novatech_dump.sql').open('w') as f:
    f.write('-- Fictional NovaTech Solutions sample data; safe to share.\n')
    for line in conn.iterdump(): f.write(line+'\n')
conn.close()

# ---------- Documents: mixed formats, deliberately cross-document facts ----------
# Facts that evaluation questions can use, with source attribution to documents.
documents = [
('01_employee_handbook.docx','docx','Employee Handbook',[
('Purpose and working principles','NovaTech Solutions is a fictional mid-sized technology services company used for this demonstration dataset. Its working principles are customer empathy, careful handling of data, written decisions, and measurable outcomes. Employees should document decisions that affect a customer commitment or a production system.'),
('Working hours and remote work','The standard collaboration window is 10:00 a.m. to 4:00 p.m. India Standard Time, Monday through Friday. Teams may agree on different schedules when customer coverage requires it. Employees working remotely must use the approved VPN when accessing internal systems. A manager should approve a recurring schedule change before it takes effect.'),
('Learning allowance','Each full-time employee may request up to INR 18,000 per financial year for job-related courses, books, or certification exams. The employee must submit the course details and business rationale to their manager before purchase. Reimbursement requires proof of payment and evidence of completion. The allowance does not roll over to the next financial year.'),
('Leave and handover','Planned leave should be submitted at least five working days in advance when practical. The employee should identify a handover contact for time-sensitive work. Sick leave should be reported to the manager as soon as reasonably possible; the five-day advance rule does not apply to unexpected illness.'),
('Information handling','Company documents should be stored in approved workspaces. Do not place customer identifiers, passwords, API keys, or private employee information in public repositories. Use fictional or anonymized records in demonstrations. Report suspected exposure of sensitive data to Security Operations promptly.')]),
('02_security_policy.pdf','pdf','Security and Access Policy',[
('Authentication','Internal services must use multi-factor authentication (MFA) wherever it is supported. Employees must never share passwords or approve an MFA prompt they did not initiate. If a device used for authentication is lost, the employee should contact the support team and their manager.'),
('Access requests','Access must follow least privilege: users receive only the permissions needed for their assigned duties. A request should name the system, requested role, business reason, and expected duration. The system owner approves access to production data. Temporary elevated access should expire at the end of the approved period.'),
('Incident reporting','A suspected security incident should be reported to Security Operations as soon as possible, ideally within one hour of discovery. The report should include the time noticed, affected service, observed behaviour, and actions already taken. Employees should preserve relevant logs and avoid deleting evidence.'),
('Credential storage','Secrets belong in the approved secret manager or deployment environment settings. They must not be committed to source control, placed in screenshots, or pasted into public issue trackers. Rotate a credential if it may have been exposed. The security team records the incident and tracks the remediation.'),
('Device requirements','Devices that access internal tools should use a supported operating system, screen lock, and current security updates. Unattended devices must be locked. Unapproved removable storage should not be used for company information.')]),
('03_support_playbook.txt','txt','Support Operations Playbook',[
('Ticket priority definitions','Critical tickets indicate a suspected security incident or a production service unavailable for many users. High tickets block a key business workflow for a team. Medium tickets affect an individual or have a workaround. Low tickets are routine requests or minor issues with no material business interruption.'),
('Response targets','The initial response target is 30 minutes for Critical, 2 business hours for High, 1 business day for Medium, and 3 business days for Low. These are response targets, not guaranteed resolution times. The support team records the first human response separately from automated acknowledgements.'),
('Escalation','Escalate a Critical ticket immediately to the on-call lead and Security Operations when security may be involved. Escalate a High ticket to the team lead if no owner has accepted it within one business hour. Include the impact, affected users, troubleshooting steps, and relevant ticket identifier.'),
('Ticket closure','A ticket can be marked Resolved after the fix or workaround is communicated to the requester. Mark it Closed after requester confirmation or after two business days without a response. Add a concise resolution note. Do not include passwords, API keys, or unnecessary personal information in ticket notes.'),
('Support analytics','Weekly reporting should include ticket volume by category, priority mix, time to first human response, resolution time for resolved tickets, reopened tickets, and the oldest open items. Compare similar periods and state the sample size before drawing conclusions.')]),
('04_data_platform_manual.docx','docx','Data Platform Manual',[
('Approved analytics workspace','The analytics workspace is the approved place to publish internal dashboards and scheduled data extracts. Dataset owners are responsible for definitions, refresh schedules, and access review. Analysts should link dashboards to a short description of the metric and its source table.'),
('Data refresh schedule','The customer activity warehouse refreshes at 06:30 IST each day. The sales summary refreshes at 07:00 IST. A dashboard should display its last successful refresh time. If a scheduled refresh fails twice in a row, the data owner should open a support ticket and inform dashboard users if the delay affects a decision.'),
('Quality checks','A routine quality check compares row counts with the previous comparable load, checks required fields for nulls, and validates known key relationships. Large changes should be investigated rather than automatically removed. Corrections should be recorded with the reason, date, and responsible owner.'),
('Metric definitions','Active customer means a customer with at least one qualifying activity event in the last 90 days. Monthly recurring revenue is the normalized monthly value of active subscriptions and excludes one-time implementation fees. A report must state the date range and timezone used for aggregation.'),
('Responsible data use','Use only fields needed for the analysis. Prefer aggregate data for broad reporting. Access to row-level customer data requires an approved business reason. When exporting results, check that the destination and intended audience are appropriate.')]),
('05_project_delivery_guide.pdf','pdf','Project Delivery Guide',[
('Project stages','A standard project uses discovery, design, implementation, validation, and handover stages. Each stage has an owner and an exit check. A project lead records scope changes and explains their impact on delivery date, cost, and quality.'),
('Change management','Changes affecting agreed scope must be written down with the reason, expected benefit, effort estimate, and risk. The project sponsor approves changes that increase the approved budget or move the target date by more than ten business days. Smaller changes may be approved by the project lead and recorded in the project log.'),
('Release readiness','Before a release, the team confirms tests have passed, rollback steps are documented, monitoring is available, and the support owner has received a handover. A release with an unresolved Critical defect must not proceed without a documented exception from the engineering lead and business owner.'),
('Project status reporting','Weekly status reports include progress against milestones, current risks, open decisions, spend against budget, and the next planned milestone. Status should be based on evidence rather than optimistic estimates. Blockers need an owner and a next action.'),
('Handover','The handover package includes architecture notes, configuration requirements, known limitations, operational contacts, and a short runbook. Secrets must be transferred using the approved secret manager rather than written into the handover document.')]),
('06_customer_service_faq.txt','txt','Customer Service FAQ',[
('Business hours','The standard customer support coverage window is 09:00 to 18:00 IST on business days. Critical production incidents are routed to the on-call process outside those hours. Public holidays may affect non-critical response times.'),
('Requesting a report','A report request should specify the business question, desired fields, date range, delivery format, and deadline. The analytics team confirms feasibility and identifies the data owner before promising a delivery date. Requests involving personal data need an approved purpose.'),
('Dashboard discrepancy','When a dashboard does not match another report, capture both report names, filter settings, date ranges, timezone, and refresh timestamps. Do not assume one report is correct before checking definitions and refresh status. Assign the investigation to the owner of the metric.'),
('Escalating a customer issue','Escalate a customer issue when a published commitment is at risk, several customers are affected, or the impact is not understood. Provide a short factual summary, the customer impact, actions tried, and the next update time. Avoid speculation in customer-facing messages.'),
('Feedback and service improvement','Support leads review recurring ticket themes monthly. Improvements should have an owner, a success measure, and a review date. A lower ticket count alone does not prove better service; changes in customer volume and reporting practices should also be considered.')])
]

styles=getSampleStyleSheet()
styles.add(ParagraphStyle(name='DocTitle', parent=styles['Title'], fontSize=18, leading=22, textColor=colors.HexColor('#17365D'), spaceAfter=16))
styles.add(ParagraphStyle(name='SectionHeading', parent=styles['Heading2'], fontSize=12, leading=15, textColor=colors.HexColor('#1F4E79'), spaceBefore=8, spaceAfter=5))
styles.add(ParagraphStyle(name='BodyCustom', parent=styles['BodyText'], fontSize=9.5, leading=13, spaceAfter=7))

for filename,fmt,title,sections in documents:
    path=DOCS/filename
    if fmt=='txt':
        out=[title, '='*len(title), 'NovaTech Solutions | Fictional sample document created for a technical assignment.','']
        for h,p in sections: out += [h, '-'*len(h), p, '']
        # Add operational examples to give meaningful depth without changing core facts.
        out += ['Working examples','-----------------',
                'Example: When a report refresh is late, the analyst checks the last successful refresh timestamp, confirms the source owner, and communicates the impact before changing any numbers.',
                'Example: When a request lacks a business reason, the owner asks for clarification instead of granting broad access or making an unsupported promise.',
                'These examples are illustrative and do not override the policy statements above.']
        path.write_text('\n'.join(out))
    elif fmt=='docx':
        doc=Document(); doc.add_heading(title,0)
        doc.add_paragraph('NovaTech Solutions | Fictional sample document created for a technical assignment.')
        for idx,(h,p) in enumerate(sections):
            if idx > 0: doc.add_page_break()
            doc.add_heading(h,level=1); doc.add_paragraph(p)
        doc.add_heading('Operational examples',level=1)
        doc.add_paragraph('Example: Before publishing a dashboard, the analyst confirms the metric definition, date range, timezone, and last successful refresh timestamp.')
        doc.add_paragraph('Example: When a request is missing an owner or business reason, the team records the gap and requests clarification before proceeding.')
        doc.add_paragraph('These examples are illustrative and do not override the policy statements above.')
        doc.save(path)
    else:
        story=[Paragraph(title,styles['DocTitle']),Paragraph('NovaTech Solutions | Fictional sample document created for a technical assignment.',styles['BodyCustom']),Spacer(1,8)]
        for idx,(h,p) in enumerate(sections):
            if idx > 0: story.append(PageBreak())
            story.append(Paragraph(h,styles['SectionHeading'])); story.append(Paragraph(p,styles['BodyCustom']))
        story.append(Paragraph('Operational examples',styles['SectionHeading']))
        story.append(Paragraph('Example: Record the relevant timestamp, owner, observed impact, and next action before escalating a service issue.',styles['BodyCustom']))
        story.append(Paragraph('Example: If the evidence is incomplete, state what is known and what remains unverified rather than guessing.',styles['BodyCustom']))
        story.append(Paragraph('These examples are illustrative and do not override the policy statements above.',styles['BodyCustom']))
        SimpleDocTemplate(str(path),pagesize=letter,rightMargin=55,leftMargin=55,topMargin=55,bottomMargin=55).build(story)

# Add a deterministic README for data assets
(DATA/'README.md').write_text('''# Sample data\n\nAll data in this folder is fictional and was created for the AnalyticsLiv RAG assignment.\n\n- `novatech.db`: ready-to-use SQLite database.\n- `schema.sql`: database schema.\n- `novatech_dump.sql`: schema and insert statements to recreate the database.\n- `documents/`: six mixed-format business documents (PDF, DOCX, TXT).\n\nThe database contains 60 projects, 100 employees, and 150 support tickets. Foreign keys connect employees to projects and support tickets to employees.\n''')
print('Created database and',len(documents),'documents in',ROOT)
with sqlite3.connect(db_path) as c:
    for table in ['projects','employees','support_tickets']:
        print(table, c.execute(f'SELECT COUNT(*) FROM {table}').fetchone()[0])
